# 客户「特斯拉投放数据源」Excel → 日合计 JSON 导出器 (本地运行, 需 openpyxl)
# 用法: py -X utf8 tesla-adsource-export.py <客户xlsx> <输出json>
# 输出交给服务器 data-import/tesla-adsource-import.cjs 回补 KoxCampaignDailyStat 历史缺口
# 列映射: 时间(7) 消费(10) 展现量(11) 点击量(12) 互动量(13) 私信进线数(14) 私信开口数(15) 私信留资数(16) [1-based]
import sys, json, datetime
import openpyxl

def num(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0

def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    next(rows)  # header
    agg = {}
    total = 0
    for r in rows:
        d = r[6]
        if not isinstance(d, datetime.datetime):
            continue
        k = d.date().isoformat()
        a = agg.setdefault(k, dict(rows=0, fee=0.0, impression=0, click=0, interaction=0, msgInquiries=0, msgOpenings=0, msgLeads=0))
        a['rows'] += 1
        a['fee'] += num(r[9])
        a['impression'] += int(num(r[10]))
        a['click'] += int(num(r[11]))
        a['interaction'] += int(num(r[12]))
        a['msgInquiries'] += int(num(r[13]))
        a['msgOpenings'] += int(num(r[14]))
        a['msgLeads'] += int(num(r[15]))
        total += 1
    days = [dict(day=k, **agg[k]) for k in sorted(agg)]
    json.dump(dict(source_file=src, generated_at=datetime.datetime.now().isoformat(timespec='seconds'), note_rows=total, days=days),
              open(dst, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'exported {len(days)} days, {total} note rows -> {dst}')
    print(f'range: {days[0]["day"]} ~ {days[-1]["day"]}, fee total = {sum(d["fee"] for d in days):.2f}')

if __name__ == '__main__':
    main()
